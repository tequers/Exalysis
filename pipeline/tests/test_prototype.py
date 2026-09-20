"""Exercise the production two-stage CLI with synthetic model responses only."""

from contextlib import redirect_stdout, redirect_stderr
import io
import json
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import Mock, patch

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import pipeline as app
from exam_roi import prototype
from exam_roi.llm import RequestLimits
from exam_roi.storage import CourseStore
from openpyxl import load_workbook


def fake_clients(*, marks=None, topic="Legal reasoning", known=False):
    marks = marks if marks is not None else [1, 1, 1, 1, 1, 5, 5, 10, 15, 30, 30]
    questions = [{"q_id": f"Q{i}", "text": f"Explain synthetic legal scenario {i}.",
                  "marks": value, "format": "mcq" if i <= 5 else "explain_derive"}
                 for i, value in enumerate(marks, 1)]
    tags = {q["q_id"]: {"topics": [topic], "quote": q["text"],
                       "rationale": "The task tests this topic.", "uncertainties": []}
            for q in questions}
    judgment = {
        "question_difficulty": [{"q_id": q["q_id"], "level": 3, "quote": q["text"],
                                 "rationale": "Apply a standard method.", "uncertainties": []}
                                for q in questions],
        "difficulty_rationale": "Apply a standard method.", "assumed_prerequisites": [],
        "prerequisite_evidence": [], "Conn": 1, "connection_rationale": "No downstream use.",
        "connection_evidence": [{"q_id": q["q_id"], "quote": q["text"]} for q in questions],
        "connection_edges": [], "unlocks": [], "prerequisites": [], "uncertainties": [],
    }
    extraction = Mock(return_value=json.dumps({"exam_year": 2025, "questions": questions}))
    analysis = Mock(side_effect=[json.dumps({"tags": tags, "new_topic_names": [] if known else [topic]}),
                                json.dumps({topic: judgment})])
    for client in (extraction, analysis):
        client.limits = RequestLimits()
        client.provider = "offline-test"
        client.model = "synthetic"
        client.reasoning_effort = None
    return {"extraction_client": extraction, "analysis_client": analysis}


class PrototypeTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.course = app.setup_course_folder(self.root / "course")
        self.paper = self.root / "2025_june.txt"
        self.paper.write_text("Synthetic compulsory exam, 100 marks.\n0 1 Explain a legal scenario.",
                              encoding="utf-8")
        self.output = io.StringIO()

    def cli(self, *args, clients=None):
        with (patch.object(sys, "argv", ["pipeline.py", str(self.course.folder), *map(str, args)]),
              patch.object(app, "configure_model_clients", return_value=clients) as config,
              patch.object(app, "load_dotenv"),
              redirect_stdout(self.output), redirect_stderr(self.output)):
            result = app.main()
        if clients is None:
            config.assert_not_called()
        return result

    def analyze(self, clients=None, *extra):
        return self.cli("add-exam", self.paper, "--prototype", "--total-marks", "100", *extra,
                        clients=clients or fake_clients())

    def rows(self):
        return json.loads((self.course.folder / "prototype/Exam_ROI_Pipeline.json").read_text("utf-8"))

    def test_full_command_and_offline_rebuild_produce_matching_reports(self):
        with CourseStore(self.course) as store:
            store.commit(taxonomy={"topics": {}})
        before = self.course.taxonomy_file.read_bytes()
        clients = fake_clients()
        self.assertEqual(self.analyze(clients), app.EXIT_OK, self.output.getvalue())
        self.assertIn("single printed mark allocation", clients["extraction_client"].call_args.args[1])
        rows = self.rows()
        self.assertEqual(rows[0]["report_status"], "unreviewed-prototype")
        self.assertEqual(rows[0]["per_exam"]["2025_june"]["marks_total"], 100)
        with CourseStore(self.course) as store:
            snapshot = store.load()
        self.assertEqual(snapshot.papers, {})
        questions = snapshot.candidates["2025_june"]["questions"]
        self.assertEqual([q["q_id"] for q in questions], [f"Q{i}" for i in range(1, 12)])
        self.assertEqual(sum(q["marks"] for q in questions), 100)
        self.assertEqual(self.course.taxonomy_file.read_bytes(), before)
        self.assertFalse(self.course.output_xlsx.exists())
        self.assertFalse(self.course.output_json.exists())
        wb = load_workbook(self.course.folder / "prototype/Exam_ROI_Pipeline.xlsx", data_only=True)
        self.addCleanup(wb.close)
        self.assertEqual(wb.sheetnames, ["ROI Scores", "Taxonomy", "Exam Log"])
        for sheet in wb:
            self.assertIn("UNREVIEWED PROTOTYPE", sheet["A1"].value)
            self.assertIn(rows[0]["generation_id"], sheet["A1"].value)
        self.assertEqual(wb["ROI Scores"]["C4"].value, rows[0]["topic"])
        self.assertEqual(wb["ROI Scores"]["I4"].value, rows[0]["priority"])
        self.assertEqual(wb["Exam Log"]["D3"].value, 100)
        wb.close()
        self.assertEqual(self.cli("rebuild", "--prototype"), app.EXIT_OK)
        rebuilt = self.rows()
        self.assertNotEqual(rows[0].pop("generation_id"), rebuilt[0].pop("generation_id"))
        self.assertEqual(rows, rebuilt)

    def test_missing_or_unmatched_stage2_quotes_do_not_block_prototype_exports(self):
        clients = fake_clients()
        tag_raw, score_raw = list(clients["analysis_client"].side_effect)
        tags = json.loads(tag_raw)
        scores = json.loads(score_raw)
        tags["tags"]["Q1"].pop("quote")
        tags["tags"]["Q2"]["quote"] = "not present in the question"
        judgment = scores["Legal reasoning"]
        judgment["question_difficulty"][0].pop("quote")
        judgment["question_difficulty"][1]["quote"] = "not present in the question"
        judgment["connection_evidence"][0].pop("quote")
        judgment["connection_evidence"][1]["quote"] = "not present in the question"
        clients["analysis_client"].side_effect = [json.dumps(tags), json.dumps(scores)]

        self.assertEqual(self.analyze(clients), app.EXIT_OK, self.output.getvalue())
        self.assertTrue((self.course.folder / "prototype/Exam_ROI_Pipeline.xlsx").exists())
        self.assertTrue((self.course.folder / "prototype/Exam_ROI_Pipeline.json").exists())
        with CourseStore(self.course) as store:
            candidate = store.load().candidates["2025_june"]
        self.assertEqual(candidate["evaluation_context"]["quote_validation"], "advisory")
        self.assertEqual(candidate["questions"][0]["topic_tagging"]["quote"], "")
        self.assertEqual(candidate["questions"][1]["topic_tagging"]["quote"],
                         "not present in the question")
        self.assertEqual(
            candidate["topic_judgments"]["Legal reasoning"]["question_difficulty"][0]["quote"], "")
        self.assertEqual(
            candidate["topic_judgments"]["Legal reasoning"]["connection_evidence"][0]["quote"], "")
        self.assertEqual(self.cli("rebuild", "--prototype"), app.EXIT_OK, self.output.getvalue())

    def test_invalid_total_stops_before_client_configuration(self):
        for value in (None, "nan", "inf", "0", "-1"):
            with self.subTest(value=value):
                args = [] if value is None else ["--total-marks", value]
                self.assertEqual(self.cli("add-exam", self.paper, "--prototype", *args),
                                 app.EXIT_SETUP_ERROR)

    def test_bad_question_marks_stop_before_stage_two(self):
        for marks in ([99], [100, 100], [None], [-1, 101]):
            with self.subTest(marks=marks):
                clients = fake_clients(marks=marks)
                self.assertEqual(self.analyze(clients), app.EXIT_PAPER_FAILURE, self.output.getvalue())
                clients["analysis_client"].assert_not_called()
                self.assertFalse((self.course.candidates_dir / "2025_june.json").exists())

    def test_later_paper_reuses_topics_and_force_retracts_old_evidence(self):
        self.assertEqual(self.analyze(), app.EXIT_OK)
        second = self.root / "2024_june.txt"
        second.write_text("Another synthetic exam.", encoding="utf-8")
        clients = fake_clients(known=True)
        self.assertEqual(self.cli("add-exam", second, "--prototype", "--total-marks", "100",
                                  clients=clients), app.EXIT_OK, self.output.getvalue())
        self.assertIn("Legal reasoning", clients["analysis_client"].call_args_list[0].args[1])
        self.assertEqual(self.rows()[0]["appearances"], 2)
        self.assertEqual(self.analyze(fake_clients(topic="Procedure"), "--force"), app.EXIT_OK)
        rows = {row["topic"]: row for row in self.rows()}
        self.assertEqual(set(rows), {"Legal reasoning", "Procedure"})
        self.assertEqual(rows["Legal reasoning"]["per_exam"]["2025_june"]["present"], 0)
        self.assertEqual(rows["Procedure"]["per_exam"]["2025_june"]["marks_total"], 100)

    def test_partial_batch_keeps_candidate_but_does_not_export(self):
        bad = self.root / "empty.txt"
        bad.write_text("", encoding="utf-8")
        result = self.cli("add-exam", self.paper, bad, "--prototype", "--total-marks", "100",
                          clients=fake_clients())
        self.assertEqual(result, app.EXIT_PAPER_FAILURE)
        self.assertTrue((self.course.candidates_dir / "2025_june.json").exists())
        self.assertFalse((self.course.folder / "prototype").exists())
        self.assertEqual(self.cli("rebuild", "--prototype"), app.EXIT_OK)

    def test_empty_prototype_rebuild_fails_without_models(self):
        self.assertEqual(self.cli("rebuild", "--prototype"), app.EXIT_EXPORT_FAILURE)
        self.assertIn("No candidates", self.output.getvalue())

    def test_export_failure_keeps_candidate_and_rebuild_recovers(self):
        with patch.object(prototype, "write_xlsx", side_effect=PermissionError("Workbook locked")):
            self.assertEqual(self.analyze(), app.EXIT_EXPORT_FAILURE)
        self.assertTrue((self.course.candidates_dir / "2025_june.json").exists())
        self.assertEqual(self.cli("rebuild", "--prototype"), app.EXIT_OK)

    def test_locked_second_destination_rolls_back_first(self):
        self.assertEqual(self.analyze(), app.EXIT_OK)
        folder = self.course.folder / "prototype"
        before = {p: p.read_bytes() for p in folder.iterdir()}
        replace = prototype.os.replace

        def locked_json(source, destination):
            if Path(destination).suffix == ".json":
                raise PermissionError("JSON locked")
            return replace(source, destination)

        with patch.object(prototype.os, "replace", side_effect=locked_json):
            self.assertEqual(self.cli("rebuild", "--prototype"), app.EXIT_EXPORT_FAILURE)
        self.assertEqual({p: p.read_bytes() for p in folder.iterdir()}, before)
        self.assertEqual(self.cli("rebuild", "--prototype"), app.EXIT_OK)

    def test_inconsistent_candidate_is_not_silently_exported(self):
        self.assertEqual(self.analyze(), app.EXIT_OK)
        path = self.course.candidates_dir / "2025_june.json"
        candidate = json.loads(path.read_text("utf-8"))
        candidate["questions"][0]["marks"] = 50
        path.write_text(json.dumps(candidate), encoding="utf-8")
        self.assertEqual(self.cli("rebuild", "--prototype"), app.EXIT_EXPORT_FAILURE)
        self.assertIn("expected 100", self.output.getvalue())

    def test_prototype_dry_run_needs_no_total_or_model_configuration(self):
        self.assertEqual(self.cli("add-exam", self.paper, "--prototype", "--dry-run"), app.EXIT_OK)
        self.assertFalse(self.course.candidates_dir.exists())

    def test_impossible_saved_topic_totals_are_rejected(self):
        self.assertEqual(self.analyze(), app.EXIT_OK)
        old_rows = self.rows()
        path = self.course.candidates_dir / "2025_june.json"
        candidate = json.loads(path.read_text("utf-8"))
        candidate["per_topic"]["Legal reasoning"].update(marks_total=900, mark_fraction=9)
        path.write_text(json.dumps(candidate), encoding="utf-8")
        self.assertEqual(self.cli("rebuild", "--prototype"), app.EXIT_EXPORT_FAILURE)
        self.assertIn("inconsistent saved topic totals", self.output.getvalue())
        self.assertEqual(self.rows(), old_rows)


if __name__ == "__main__":
    unittest.main()
