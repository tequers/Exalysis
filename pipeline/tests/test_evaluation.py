"""Offline checks for ticket 1; no API credentials or model calls required."""
import copy
import hashlib
import importlib
import json
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch

PIPELINE = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PIPELINE))
from exam_roi.evaluation import (
    CONTRACT_VERSION, CONTRACT_SHA256, CONTRACT_TEXT, LEGACY_VERSION,
    contract_metadata, difficulty_version, summarize_difficulty, validate_topic_scores,
)
with patch.dict(os.environ, {"LLM_PROVIDER": "anthropic"}):
    app = importlib.import_module("pipeline")

FIXTURES = json.loads((Path(__file__).parent / "fixtures/difficulty-v1.json").read_text())


def question(example, fmt="short_answer", marks=10):
    return {"q_id": "Q1", "text": example["text"], "topics": ["Topic"],
            "format": fmt, "marks": marks}


def score(example):
    return {"Topic": {
        "question_difficulty": [{"q_id": "Q1", "level": example["level"],
                                 "quote": example["quote"], "rationale": example["rationale"],
                                 "uncertainties": []}],
        "difficulty_rationale": example["rationale"],
        "assumed_prerequisites": example["assumed_prerequisites"],
        "prerequisite_evidence": [
            {"prerequisite": name, "q_id": "Q1", "quote": example["quote"],
             "rationale": "The task uses this background: " + name}
            for name in example["assumed_prerequisites"]],
        "Conn": 1, "connection_rationale": "No direct downstream use is shown in Q1.",
        "connection_evidence": [{"q_id": "Q1", "quote": example["quote"]}],
        "connection_edges": [],
        "unlocks": [], "prerequisites": [], "uncertainties": []}}


def tagging(topic, quote):
    return {"Q1": {
        "topics": [topic], "quote": quote,
        "rationale": f"The question directly tests {topic}.", "uncertainties": [],
    }}


class ContractTests(unittest.TestCase):
    def test_contract_is_versioned_and_prompts_share_it(self):
        self.assertIn("# Evaluation contract " + CONTRACT_VERSION, CONTRACT_TEXT)
        self.assertIn(CONTRACT_TEXT, app._S1_SYSTEM)
        self.assertIn(CONTRACT_TEXT, app._S2_SYSTEM)
        # Preserve earlier prompt text as well as the current contract snapshot.
        previous = (PIPELINE / "exam_roi/contracts/evaluation-v1.md").read_text(encoding="utf-8")
        self.assertEqual(hashlib.sha256(previous.encode()).hexdigest(),
                         "f48b4d83666f56bfb33130aeae7c4f266252ba74c9ff5f032453a9eebf21394b")
        previous_wording = (PIPELINE / "exam_roi/contracts/evaluation-v1.0.1.md").read_text(encoding="utf-8")
        self.assertEqual(hashlib.sha256(previous_wording.encode()).hexdigest(),
                         "98d3f6cadaf41fb615fdb69574799c742a1cbaa5df9003737abee4b8648028c0")
        self.assertEqual(CONTRACT_SHA256,
                         "5674324e8a2ca455c606279fe28d166b43498767b4793edc8d7cfdcee877f889")
        self.assertNotIn("Diff_hours", app._S2B_TEMPLATE)
        self.assertNotIn("exam marks per hour", app.HELP)

    def test_all_anchors_and_adjacent_boundaries_have_reviewed_examples(self):
        examples = {e["id"]: e for e in FIXTURES["examples"]}
        self.assertEqual({e["level"] for e in examples.values()}, set(range(1, 7)))
        self.assertEqual(FIXTURES["review"]["status"], "author-reviewed")
        self.assertFalse(FIXTURES["review"]["human_reviewed"])
        pairs = set()
        for boundary in FIXTURES["boundaries"]:
            low, high = examples[boundary["lower"]], examples[boundary["upper"]]
            pairs.add((low["level"], high["level"]))
            self.assertTrue(boundary["distinction"])
        self.assertEqual(pairs, {(n, n + 1) for n in range(1, 6)})
        for example in examples.values():
            with self.subTest(example=example["id"]):
                self.assertIn(f"| {example['level']} |", CONTRACT_TEXT)
                self.assertTrue(example["counterexample"])
                out = validate_topic_scores(score(example), ["Topic"], [question(example)],
                                          ["Topic"])
                self.assertEqual(out["Topic"]["Diff"], example["level"])
                self.assertEqual(out["Topic"]["evaluation_contract_version"], CONTRACT_VERSION)

    def test_ordinal_median_uses_equal_paper_weights_and_retains_conflict(self):
        result = summarize_difficulty({"paper_a": [1] * 20, "paper_b": [5], "paper_c": [5]})
        self.assertEqual(result["Diff"], 5)
        self.assertEqual(result["difficulty_range"], [1, 5])
        self.assertTrue(result["difficulty_review_required"])
        self.assertFalse(result["difficulty_provisional"])
        self.assertEqual(summarize_difficulty({"a": [2, 5]})["Diff"], 2)
        self.assertEqual(summarize_difficulty({"a": [2], "b": [3]})["Diff"], 2)
        self.assertFalse(summarize_difficulty({"a": [2, 3]})["difficulty_review_required"])

    def test_unresolved_and_invalid_levels_do_not_become_defaults(self):
        for invalid in ({}, {"a": []}, {"a": [None]}, {"a": [True]},
                        {"a": [0]}, {"a": [7]}, {"a": [2.5]}, {"a": ["3"]}):
            with self.subTest(invalid=invalid), self.assertRaises(ValueError):
                summarize_difficulty(invalid)

    def test_difficulty_inputs_and_outputs_are_independent_of_format_and_marks(self):
        examples = {e["id"]: e for e in FIXTURES["examples"]}
        for pair in FIXTURES["format_pairs"]:
            example = examples[pair["example"]]
            outputs = []
            for fmt, marks in zip(pair["formats"], [1, 100]):
                with patch.object(app, "call_llm", return_value=json.dumps(score(example))) as call:
                    out = app._stage2_score_topics(["Topic"], "", [question(example, fmt, marks)],
                                               ["Topic"], client=app.call_llm)
                prompt = call.call_args.args[1]
                self.assertIn(example["text"], prompt)
                self.assertNotIn('"format":', prompt)
                self.assertNotIn('"marks":', prompt)
                self.assertEqual(out["Topic"]["Diff"], pair["expected_level"])
                outputs.append(out)
            self.assertEqual(outputs[0], outputs[1])

    def test_judgment_and_inferred_background_evidence_must_be_resolvable(self):
        example = FIXTURES["examples"][2]
        modifications = [
            lambda s: s["Topic"]["question_difficulty"][0].update(quote="invented quotation"),
            lambda s: s["Topic"]["question_difficulty"][0].update(q_id="unknown"),
            lambda s: s["Topic"]["question_difficulty"][0].update(rationale=""),
            lambda s: s["Topic"].update(assumed_prerequisites=[]),
            lambda s: s["Topic"].update(prerequisite_evidence=[]),
            lambda s: s["Topic"]["prerequisite_evidence"][0].update(quote="invented"),
            lambda s: s["Topic"]["prerequisite_evidence"][0].update(q_id="unknown"),
            lambda s: s["Topic"]["prerequisite_evidence"][0].update(rationale=""),
            lambda s: s["Topic"].update(prerequisites=["Absent topic"]),
            lambda s: s["Topic"].update(unlocks=["Topic"]),
            lambda s: s["Topic"].update(Conn=True),
            lambda s: s["Topic"].update(connection_evidence=[]),
            lambda s: s["Topic"].update(question_difficulty=[]),
        ]
        for mutate in modifications:
            s = score(example)
            mutate(s)
            with self.subTest(data=s), self.assertRaises(ValueError):
                validate_topic_scores(s, ["Topic"], [question(example)], ["Topic"])

    def test_unrequested_duration_fields_are_rejected(self):
        example = FIXTURES["examples"][0]
        s = score(example)
        s["Topic"].update(Diff_hours="99h", difficulty_hours=99)
        with self.assertRaisesRegex(ValueError, "unsupported: Diff_hours, difficulty_hours"):
            validate_topic_scores(s, ["Topic"], [question(example)], ["Topic"])

    def test_empty_inferred_background_is_valid_without_manual_input(self):
        example = FIXTURES["examples"][0]
        answer = score(example)
        answer["Topic"].update(assumed_prerequisites=[], prerequisite_evidence=[])
        out = validate_topic_scores(answer, ["Topic"], [question(example)], ["Topic"])
        self.assertEqual(out["Topic"]["assumed_prerequisites"], [])
        self.assertEqual(out["Topic"]["Diff"], 1)

    def test_legacy_manual_prerequisites_do_not_affect_scoring(self):
        example = FIXTURES["examples"][2]
        outputs, prompts = [], []
        for taxonomy in ({"topics": {}}, {"topics": {}, "course_prerequisites":
                                         {"invalid_legacy_value": "MANUAL_SENTINEL"}}):
            before = copy.deepcopy(taxonomy)
            with patch.object(app, "_stage2_tag", return_value=(tagging("Topic", example["quote"]), ["Topic"])), \
                 patch.object(app, "call_llm", return_value=json.dumps(score(example))) as call:
                outputs.append(app.stage2_tag_score([question(example)], taxonomy, 10, client=app.call_llm))
                prompts.append(call.call_args.args[1])
            self.assertEqual(taxonomy, before)
        self.assertEqual(outputs[0], outputs[1])
        self.assertEqual(prompts[0], prompts[1])
        self.assertNotIn("MANUAL_SENTINEL", prompts[0])

    def test_contract_module_import_has_no_cli_or_provider_side_effects(self):
        env = {**os.environ, "PYTHONPATH": str(PIPELINE), "LLM_PROVIDER": "invalid-on-purpose"}
        script = "import sys; before=sys.stdout; import exam_roi.evaluation; assert sys.stdout is before"
        run = subprocess.run([sys.executable, "-c", script], env=env, capture_output=True, text=True)
        self.assertEqual(run.returncode, 0, run.stderr)
        self.assertEqual(run.stdout, "")


class WorkflowTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.folder = Path(self.tmp.name)
        self.context = app.setup_course_folder(self.folder)
        self.context.parsed_dir.mkdir()
        self.example = FIXTURES["examples"][2]
        self.taxonomy = {"topics": {"Legacy": {"Diff": 4, "Conn": 2,
                         "Diff_hours": "3–6h", "difficulty_hours": 6,
                         "prerequisites": [], "first_seen": "old"}}}
        app.save_taxonomy(self.taxonomy, course=self.context)
        self.old = {"exam_id": "old", "year": 2024, "total_marks": 100,
                    "source_file": "old.txt", "questions": [], "per_topic": {
                    "Legacy": {"marks_total": 20, "mark_fraction": .2,
                    "format_distribution": {"short_answer": 1}, "Diff_hours": "3–6h"}}}
        (self.context.parsed_dir / "old.json").write_text(json.dumps(self.old), encoding="utf-8")

    def test_legacy_rebuild_preserves_records_and_removes_duration_from_both_exports(self):
        from openpyxl import load_workbook
        before = {p: p.read_bytes() for p in [self.context.taxonomy_file, self.context.parsed_dir / "old.json"]}
        app.cmd_rebuild(None, course=self.context)
        rows = json.loads(self.context.output_json.read_text(encoding="utf-8"))
        self.assertEqual(rows[0]["Diff"], 4)
        self.assertEqual(rows[0]["Conn"], 2)
        self.assertEqual(rows[0]["priority"], 5)
        self.assertEqual(rows[0]["difficulty_contract_version"], LEGACY_VERSION)
        self.assertEqual(rows[0]["per_exam"]["old"]["evaluation_contract_version"], LEGACY_VERSION)
        self.assertNotIn("hours", json.dumps(rows))
        wb = load_workbook(self.context.output_xlsx)
        try:
            cells = [str(c.value) for sheet in wb for row in sheet for c in row if c.value is not None]
            self.assertFalse(any("hour" in c.lower() or c == "3–6h" for c in cells))
            self.assertIn(LEGACY_VERSION, cells)
            self.assertIn("Diff contract", cells)
            self.assertIn("Analysis contract", cells)
            self.assertEqual(wb["ROI Scores"]["I4"].value, rows[0]["priority"])
            self.assertEqual(wb["ROI Scores"].max_column, 15)
        finally:
            wb.close()
        for path, content in before.items():
            self.assertEqual(path.read_bytes(), content)

    def _process_new(self, answer=None):
        paper = self.folder / "new_2026.txt"
        paper.write_text(self.example["text"], encoding="utf-8")
        q = question(self.example)
        q.pop("topics")
        with patch.object(app, "stage1_extract", return_value=([q], 2026)), \
             patch.object(app, "_stage2_tag", return_value=(tagging("Topic", self.example["quote"]), ["Topic"])), \
             patch.object(app, "call_llm", return_value=json.dumps(answer or score(self.example))):
            app.process_exam_file(paper, force=True, course=self.context, extraction_client=app.call_llm, analysis_client=app.call_llm)
        return json.loads((self.context.candidates_dir / "new_2026.json").read_text(encoding="utf-8"))

    def test_new_analysis_records_version_and_grounded_qualitative_judgments(self):
        parsed = self._process_new()
        self.assertEqual(parsed["evaluation_contract_sha256"], CONTRACT_SHA256)
        self.assertEqual(parsed["evaluation_contract_version"], CONTRACT_VERSION)
        judgment = parsed["per_topic"]["Topic"]
        self.assertEqual(judgment["Diff"], 3)
        self.assertNotIn("course_prerequisites", app.load_taxonomy(course=self.context))
        self.assertEqual(judgment["prerequisite_evidence"][0]["q_id"], "Q1")
        self.assertEqual(judgment["question_difficulty"][0]["quote"], self.example["quote"])
        self.assertNotIn("hours", json.dumps(parsed))
        taxonomy = app.load_taxonomy(course=self.context)
        self.assertEqual(taxonomy, self.taxonomy)
        proposal = parsed["proposed_taxonomy_changes"]["topics"]["Topic"]
        self.assertEqual(proposal["evaluation_contract_version"], CONTRACT_VERSION)
        self.assertEqual(proposal["first_seen"], "new_2026")

    def test_abstained_new_difficulty_does_not_overwrite_course_scores(self):
        before = self.context.taxonomy_file.read_bytes()
        answer = score(self.example)
        answer["Topic"]["question_difficulty"][0]["level"] = None
        with self.assertRaisesRegex(ValueError, "resolved integer"):
            self._process_new(answer)
        self.assertEqual(self.context.taxonomy_file.read_bytes(), before)
        self.assertFalse((self.context.candidates_dir / "new_2026.json").exists())

    def test_explicit_override_preserved_when_topic_reappears(self):
        from argparse import Namespace
        with patch.object(app, "cmd_rebuild"):
            app.cmd_edit_topic(Namespace(name="Legacy", diff=5, conn=None), course=self.context)
        topic = app.load_taxonomy(course=self.context)["topics"]["Legacy"]
        self.assertEqual(topic["Conn"], 2)
        self.assertNotIn("evaluation_contract_version", topic)
        self.assertEqual(difficulty_version(topic), CONTRACT_VERSION)
        self.assertEqual(topic["human_overrides"]["Diff"]["previous_value"], 4)
        q = question(self.example)
        q["topics"] = ["Legacy"]
        taxonomy = app.load_taxonomy(course=self.context)
        before = copy.deepcopy(taxonomy)
        with patch.object(app, "_stage2_tag", return_value=(tagging("Legacy", self.example["quote"]), [])), \
             patch.object(app, "call_llm", return_value=json.dumps({"Legacy": score(self.example)["Topic"]})) as call:
            result = app.stage2_tag_score([q], taxonomy, 10, client=app.call_llm)
        call.assert_called_once()
        self.assertEqual(taxonomy, before)
        self.assertEqual(result["per_topic"]["Legacy"]["Diff"], 3)
        updated = app.aggregate_taxonomy(taxonomy, {"new": {**contract_metadata(), **result}})
        self.assertEqual(updated["topics"]["Legacy"]["Diff"], 5)
        self.assertEqual(updated["topics"]["Legacy"]["model_estimate"]["Diff"], 3)
        self.assertEqual(updated["topics"]["Legacy"]["human_overrides"], topic["human_overrides"])

    def test_cli_compatible_from_root_and_pipeline_directory(self):
        env = {**os.environ, "LLM_PROVIDER": "anthropic"}
        for cwd, entry in [(PIPELINE.parent, "pipeline/pipeline.py"), (PIPELINE, "pipeline.py")]:
            for args in [[], [str(self.folder), "status"], [str(self.folder), "rebuild"]]:
                run = subprocess.run([sys.executable, entry, *args], cwd=cwd, env=env,
                                     capture_output=True, text=True, encoding="utf-8")
                self.assertEqual(run.returncode, 0, run.stderr)
                self.assertNotIn("per hour", run.stdout)


if __name__ == "__main__":
    unittest.main()
